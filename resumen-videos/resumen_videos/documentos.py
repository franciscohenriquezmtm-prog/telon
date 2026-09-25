"""Generación del manual didáctico: ``.docx`` (python-docx) y ``.pdf`` (reportlab).

Estructura de ambos documentos: portada, índice por secciones (opcional; una
columna, títulos completos y número de página de cada paso) y las páginas de
pasos en rejilla (1x1, 1x2, 1x3 en A4 vertical; 2x2 en A4 horizontal; con
capturas verticales, 2x1 en A4 horizontal).  Cada celda lleva la captura, la
línea "PASO N · SECCIÓN · mm:ss", el título en negrita y la descripción.

Los textos nunca se recortan en silencio: en cada celda la imagen se reduce lo
justo para que quepan título y descripción completos (hasta un tope de líneas),
luego se baja 1 pt la fuente y solo en último término se recorta con "…" y se
avisa con el número de paso (``avisos``).

Técnicas heredadas del prototipo verificado: en el docx las filas tienen alto
EXACTO y el ancho se fija en cada celda (paginación determinista), la tabla no
lleva estilo (sin bordes), la cabecera de cada página fuerza el salto ANTES y el
documento termina en un párrafo de 1 pt; en el PDF se dibuja sobre el canvas con
coordenadas absolutas, los textos se miden con ``Paragraph.wrap`` (con las
mismas fuentes TrueType se maquetan igual para docx y pdf) y las capturas JPEG se
incrustan sin recodificar.
"""
from __future__ import annotations

import datetime as _dt
import math
import os
import re
import statistics
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable
from xml.sax.saxutils import escape as _xml_escape

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.image.exceptions import InvalidImageStreamError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib.utils import ImageReader, simpleSplit
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import Paragraph

from . import config
from .modelos import Momento, formatear_tiempo

# ----------------------------------------------------------------------------- maqueta
A4_ANCHO_CM, A4_ALTO_CM = 21.0, 29.7
MARGEN_CM = 1.0             # margen de página (los 4 lados) en las páginas de índice y de pasos
GAP_X_CM = 0.5              # separación horizontal entre celdas
GAP_Y_CM = 0.4              # separación vertical entre filas
CABECERA_CM = 0.8           # cabecera (texto gris + línea fina) de índice y pasos
SEGURIDAD_CM = 0.3          # holgura para redondeos de Word
CELDA_MARGEN_LAT_CM = 0.15  # margen interno izquierdo/derecho de cada celda
CELDA_MARGEN_SUP_CM = 0.05  # margen interno superior de cada celda
INTERLINEADO = 1.2          # interlineado EXACTO (pt = tamaño x 1.2)
# Lo que ocupa una celda además de la imagen y las líneas de texto.  En el docx la fila es de alto EXACTO y
# Word recorta lo que sobre, así que se reserva también el descendente de la línea de la imagen y una holgura.
ESPACIO_TRAS_IMAGEN_PT = 2.0      # space_after del párrafo de la imagen
ESPACIO_ENTRE_PARRAFOS_PT = 1.0   # tras la línea de paso y tras el título
DESCENDENTE_IMAGEN_PT = 3.0       # descendente de la línea que contiene la imagen (marca de párrafo a 2 pt)
SEGURIDAD_CELDA_PT = 4.0          # holgura mínima que queda libre al pie de cada celda del docx
PAD_TEXTO_PT = ESPACIO_TRAS_IMAGEN_PT + 2 * ESPACIO_ENTRE_PARRAFOS_PT + DESCENDENTE_IMAGEN_PT + SEGURIDAD_CELDA_PT
RATIO_CAJA = 16 / 9         # aspecto de la caja de imagen si no se puede leer ninguna captura
PT_CABECERA = 8.0
PT_INDICE_TITULO = 18.0
PT_INDICE_SECCION = 11.0
PT_INDICE_PASO = 10.0
ALTO_INDICE_SECCION_PT = 15.0     # por línea
ALTO_INDICE_PASO_PT = 14.0        # por línea (un título largo ocupa varias)
ESPACIO_INDICE_SECCION_PT = 10.0
ALTO_TITULO_INDICE_PT = PT_INDICE_TITULO * INTERLINEADO + 14.0
INDICE_ANCHO_DERECHA_CM = 3.2     # columna derecha del índice: "mm:ss · pág. NNN"
MAX_LINEAS_TITULO = 3             # líneas que puede ocupar el título antes de bajar la fuente y recortar
SECCION_POR_DEFECTO = "General"
PT_TRANSCRIPCION = 9.5            # anexo de transcripción: cuerpo
PT_TRANSCRIPCION_TITULO = 18.0
TRANSCRIPCION_COL_CM = 1.5        # columna izquierda con el mm:ss de cada segmento
ESPACIO_SEGMENTO_PT = 3.0         # separación entre segmentos

COLOR_GRIS = "#666666"      # metadatos, cabeceras, línea de paso
COLOR_TITULO = "#111111"    # títulos
COLOR_TEXTO = "#222222"     # descripciones
COLOR_LINEA = "#BBBBBB"     # líneas finas
COLOR_CAJA = "#E5E5E5"      # caja "[sin captura]"

# (título, descripción, línea de paso) en pt según pasos por página
_FUENTES = {4: (10.5, 8.5, 7.5), 3: (11.5, 9.5, 8.0), 2: (13.0, 10.5, 8.5), 1: (15.0, 11.5, 9.0)}
# líneas de descripción reservadas al calcular el alto de referencia de la imagen
_LINEAS_DESC = {4: 3, 3: 2, 2: 3, 1: 4}
# tope de líneas de descripción: la imagen de la celda se reduce hasta que quepan (nunca se recorta en silencio)
_LINEAS_DESC_MAX = {4: 6, 3: 8, 2: 8, 1: 8}
_CARACTERES_PROHIBIDOS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")   # inválidos en XML (docx) y sin glifo en el PDF
# Con fuentes estándar (cp1252) estos símbolos se escriben con un equivalente legible en vez de '?'.
_SUSTITUCIONES_CP1252 = {"≤": "<=", "≥": ">=", "→": "->", "←": "<-", "−": "-", "′": "'", "″": '"',
                         "≈": "~", "≠": "!=", "∆": "delta ", "Δ": "delta ", "∞": "inf."}
# Fuentes TrueType del sistema que se usan en el PDF si falta la carpeta fuentes/ del proyecto:
# (regular, negrita, nombre con el que se registra).  Se prueba en orden y se usa la primera que exista.
_WINDIR = Path(os.environ.get("WINDIR") or os.environ.get("SystemRoot") or "C:/Windows")
FUENTES_SISTEMA = (
    (_WINDIR / "Fonts" / "arial.ttf", _WINDIR / "Fonts" / "arialbd.ttf", "Arial"),
    (_WINDIR / "Fonts" / "segoeui.ttf", _WINDIR / "Fonts" / "segoeuib.ttf", "SegoeUI"),
    (Path("/Library/Fonts/Arial.ttf"), Path("/Library/Fonts/Arial Bold.ttf"), "Arial"),
    (Path("/System/Library/Fonts/Supplemental/Arial.ttf"), Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
     "Arial"),
    (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
     "DejaVu"),
)
_aviso_fuentes_emitido = False


@dataclass(frozen=True)
class Layout:
    """Números de la maqueta para ``n`` pasos (todo en cm salvo los tamaños en pt)."""

    n: int
    por_pagina: int
    cols: int
    filas: int
    horizontal: bool
    paginas_contenido: int
    pag_w_cm: float
    pag_h_cm: float
    util_w_cm: float
    util_h_cm: float
    celda_w_cm: float
    celda_h_cm: float
    interior_w_cm: float
    img_w_cm: float
    img_h_cm: float
    pt_titulo: float
    pt_desc: float
    pt_meta: float
    pt_cabecera: float
    lineas_desc: int           # líneas de descripción que caben bajo la imagen de referencia
    lineas_desc_max: int       # tope de líneas: la imagen se reduce hasta que quepan (luego fuente -1 pt, luego recorte)
    ratio: float               # aspecto (ancho/alto) de la caja de imagen
    vertical: bool             # las capturas son verticales (iPhone en vertical): rejilla adaptada

    @property
    def descripcion(self) -> str:
        """Texto para el log: ``"A4 horizontal 2x2"`` (columnas x filas)."""
        return f"A4 {'horizontal' if self.horizontal else 'vertical'} {self.cols}x{self.filas}"


def normalizar_por_pagina(valor) -> str | int:
    """Devuelve ``"auto"`` o un entero 1-4; acepta también cadenas (``"2"``).  ValueError si no vale."""
    if isinstance(valor, str):
        texto = valor.strip().lower()
        if texto == "auto":
            return "auto"
        if not texto.isdigit():
            raise ValueError(f"por_pagina inválido: {valor!r} (use auto, 1, 2, 3 o 4)")
        valor = int(texto)
    if isinstance(valor, bool) or not isinstance(valor, int) or valor not in (1, 2, 3, 4):
        raise ValueError(f"por_pagina inválido: {valor!r} (use auto, 1, 2, 3 o 4)")
    return valor


def calcular_layout(n: int, por_pagina="auto", ratio: float = RATIO_CAJA) -> Layout:
    """Maqueta para ``n`` pasos.  ``"auto"``: n <= 5 -> 1; 6-10 -> 2; 11-15 -> 3; >= 16 -> 4 (2x2 horizontal).

    ``ratio`` es el aspecto (ancho/alto) de la caja reservada a la captura; las imágenes se ajustan
    dentro de ella conservando su propio aspecto.  Con capturas verticales (``ratio < 1``, iPhone en
    vertical) "auto" usa como máximo 2 por página, lado a lado en A4 horizontal (2x1), con la imagen a
    toda la altura que deja el texto: en 1x3/2x2 quedarían de menos de 4 cm de ancho.
    ``img_w_cm``/``img_h_cm`` son la caja de referencia (máxima): en cada celda la imagen se reduce lo
    justo para que el texto completo quepa (ver ``_maquetar_celda``).
    """
    n = max(1, int(n))
    pp = normalizar_por_pagina(por_pagina)
    ratio = min(max(float(ratio), 0.5), 2.4)
    vertical = ratio < 1.0
    if pp == "auto":
        pp = 1 if n <= 5 else 2 if (vertical or n <= 10) else 3 if n <= 15 else 4
    if pp == 4:
        cols, filas, horizontal = 2, 2, True
    elif pp == 2 and vertical:
        cols, filas, horizontal = 2, 1, True
    else:
        cols, filas, horizontal = 1, pp, False
    pag_w, pag_h = (A4_ALTO_CM, A4_ANCHO_CM) if horizontal else (A4_ANCHO_CM, A4_ALTO_CM)
    util_w = pag_w - 2 * MARGEN_CM
    util_h = pag_h - 2 * MARGEN_CM - CABECERA_CM - SEGURIDAD_CM
    celda_w = (util_w - (cols - 1) * GAP_X_CM) / cols
    celda_h = (util_h - (filas - 1) * GAP_Y_CM) / filas
    pt_titulo, pt_desc, pt_meta = _FUENTES[pp]
    lineas = _LINEAS_DESC[pp]
    texto_h_pt = (pt_meta + pt_titulo + lineas * pt_desc) * INTERLINEADO + PAD_TEXTO_PT
    texto_h = texto_h_pt / 72 * 2.54
    interior_w = celda_w - 2 * CELDA_MARGEN_LAT_CM
    img_h = min(celda_h - CELDA_MARGEN_SUP_CM - texto_h, interior_w / ratio)
    img_w = img_h * ratio
    alto_libre_pt = ((celda_h - CELDA_MARGEN_SUP_CM - img_h) / 2.54 * 72
                     - (pt_meta + pt_titulo) * INTERLINEADO - PAD_TEXTO_PT)
    lineas_reales = max(1, int(alto_libre_pt / (pt_desc * INTERLINEADO) + 1e-6))   # 2.9999999 -> 3, 2.9 -> 2
    return Layout(n=n, por_pagina=pp, cols=cols, filas=filas, horizontal=horizontal,
                  paginas_contenido=math.ceil(n / pp), pag_w_cm=pag_w, pag_h_cm=pag_h,
                  util_w_cm=util_w, util_h_cm=util_h, celda_w_cm=celda_w, celda_h_cm=celda_h,
                  interior_w_cm=interior_w, img_w_cm=img_w, img_h_cm=img_h, pt_titulo=pt_titulo,
                  pt_desc=pt_desc, pt_meta=pt_meta, pt_cabecera=PT_CABECERA, lineas_desc=lineas_reales,
                  lineas_desc_max=max(lineas_reales, _LINEAS_DESC_MAX[pp]), ratio=ratio, vertical=vertical)


# ----------------------------------------------------------------------------- utilidades
@dataclass
class _Datos:
    """Todo lo que necesitan la portada y las cabeceras."""

    nombre_video: str
    titulo: str
    resumen: str | None
    fecha: str
    duracion: float | None
    modo: str
    modelo: str
    momentos: list
    secciones: int
    transcripcion: list = field(default_factory=list)   # [{"inicio": s, "fin": s, "texto": str}] (anexo)

    @property
    def pie(self) -> str:
        partes = ["Generado con resumen_videos"]
        if self.modo:
            partes.append(f"modo {self.modo}")
        if self.modelo:
            partes.append(f"modelo {self.modelo}")
        return " · ".join(partes)

    def filas_portada(self) -> list[tuple[str, str]]:
        filas = [("Video", self.nombre_video), ("Fecha", self.fecha)]
        if self.duracion is not None:
            filas.append(("Duración", f"{formatear_tiempo(self.duracion)} (mm:ss)"))
        filas.append(("Pasos", str(len(self.momentos))))
        if self.secciones:
            filas.append(("Secciones", str(self.secciones)))
        return filas


def _paginar(items: list, k: int) -> list[list]:
    return [items[i:i + k] for i in range(0, len(items), k)]


def _limpiar(texto) -> str:
    """Quita los caracteres de control (inválidos en el XML del docx) y colapsa los espacios."""
    return " ".join(_CONTROL.sub("", str(texto or "")).split())


def _recortar(texto, max_chars: int) -> str:
    """Limpia el texto y lo recorta en un límite de palabra añadiendo '…' si hace falta."""
    texto = _limpiar(texto)
    if len(texto) <= max_chars:
        return texto
    corte = texto[:max_chars].rsplit(" ", 1)[0].rstrip(" ,;:.")
    return corte + "…"


def _nombre_archivo_seguro(nombre: str) -> str:
    """Nombre de archivo válido en Windows (sin ``: * ? " < > |`` ni punto/espacio final)."""
    limpio = _CARACTERES_PROHIBIDOS.sub("_", str(nombre or "")).strip(" .")
    return limpio or "video"


def _tamano_imagen(ruta) -> tuple[int, int] | None:
    """(ancho, alto) en píxeles leídos con Pillow, o None si no se puede abrir."""
    if not ruta:
        return None
    try:
        with Image.open(ruta) as im:
            return im.size
    except (OSError, ValueError):
        return None


def _ratio_capturas(momentos: list) -> float:
    """Aspecto mediano de las capturas legibles (la caja de imagen se adapta a él)."""
    aspectos = []
    for m in momentos:
        tam = _tamano_imagen(m.captura_para_documento)
        if tam and tam[1] > 0:
            aspectos.append(tam[0] / tam[1])
    return statistics.median(aspectos) if aspectos else RATIO_CAJA


def _grupos_indice(momentos: list) -> list[tuple[str, list[tuple[int, Momento]]]]:
    """Agrupa los pasos por sección consecutiva: [(nombre, [(numero, momento), ...]), ...]."""
    grupos: list[tuple[str, list]] = []
    for i, m in enumerate(momentos, 1):
        nombre = " ".join(str(m.seccion or "").split()) or SECCION_POR_DEFECTO
        if not grupos or grupos[-1][0] != nombre:
            grupos.append((nombre, []))
        grupos[-1][1].append((i, m))
    return grupos


def _contar_secciones(momentos: list) -> int:
    return len({nombre for nombre, _ in _grupos_indice(momentos)})


def _segmentos_transcripcion(transcripcion) -> list[tuple[str, str]]:
    """``[(mm:ss, texto)]`` limpios y en orden a partir de la transcripción guardada; entradas raras se ignoran."""
    segmentos: list[tuple[float, str]] = []
    for s in transcripcion or []:
        if not isinstance(s, dict):
            continue
        texto = _limpiar(s.get("texto"))
        inicio = s.get("inicio")
        if not texto or isinstance(inicio, bool) or not isinstance(inicio, (int, float)):
            continue
        segmentos.append((max(0.0, float(inicio)), texto))
    segmentos.sort(key=lambda s: s[0])
    return [(formatear_tiempo(t), texto) for t, texto in segmentos]


def _preparar(nombre_video, momentos, titulo, resumen, fecha, duracion, modo, modelo,
              por_pagina, log: Callable[[str], None] | None = None,
              transcripcion: list | None = None) -> tuple[_Datos, Layout, dict]:
    """Valida las entradas y calcula datos, maqueta y fuentes (común a docx y pdf)."""
    momentos = list(momentos or [])
    if not momentos:
        raise ValueError("no hay momentos: no se puede generar el documento")
    datos = _Datos(nombre_video=_limpiar(nombre_video) or "video",
                   titulo=_recortar(titulo or nombre_video or "video", 120) or "video",
                   resumen=_recortar(resumen, 700) or None,
                   fecha=_limpiar(fecha) or _dt.date.today().strftime("%d/%m/%Y"),
                   duracion=duracion, modo=_limpiar(modo), modelo=_limpiar(modelo),
                   momentos=momentos, secciones=_contar_secciones(momentos),
                   transcripcion=_segmentos_transcripcion(transcripcion))
    layout = calcular_layout(len(momentos), por_pagina, ratio=_ratio_capturas(momentos))
    return datos, layout, _fuentes_pdf(log)


# ----------------------------------------------------------------------------- fuentes y medidas (PDF)
def _candidatas_fuentes() -> list[tuple[Path, Path, str]]:
    """Fuentes TrueType a probar para el PDF, en orden: las del proyecto (``fuentes/``) y luego las del sistema."""
    carpeta = Path(config.CARPETA_FUENTES)
    return [(carpeta / "DejaVuSans.ttf", carpeta / "DejaVuSans-Bold.ttf", "DejaVu"), *FUENTES_SISTEMA]


def _fuentes_pdf(log: Callable[[str], None] | None = None) -> dict:
    """Registra la primera fuente TrueType disponible (Unicode) para el PDF.

    Orden: DejaVu de ``config.CARPETA_FUENTES`` (incluida en el proyecto), luego las del sistema
    (``FUENTES_SISTEMA``: Arial/Segoe UI en Windows, Arial en macOS, DejaVu en Linux).  Solo si no hay
    ninguna se usa Helvetica con cp1252 (``latin1``: los símbolos fuera de cp1252 se sustituyen), y se avisa
    una vez por ``log``.  Devuelve ``{"regular", "negrita", "latin1", "origen"}``.
    """
    global _aviso_fuentes_emitido
    for regular, negrita, nombre in _candidatas_fuentes():
        if not Path(regular).is_file():
            continue
        registradas = pdfmetrics.getRegisteredFontNames()
        if nombre not in registradas:
            pdfmetrics.registerFont(TTFont(nombre, str(regular)))
        nombre_negrita = nombre
        if Path(negrita).is_file():
            nombre_negrita = f"{nombre}-Bold"
            if nombre_negrita not in registradas:
                pdfmetrics.registerFont(TTFont(nombre_negrita, str(negrita)))
        return {"regular": nombre, "negrita": nombre_negrita, "latin1": False, "origen": str(regular)}
    if log is not None and not _aviso_fuentes_emitido:
        _aviso_fuentes_emitido = True
        log(f"  aviso: no se encontró ninguna fuente TrueType ({Path(config.CARPETA_FUENTES) / 'DejaVuSans.ttf'} ni las "
            "del sistema); el PDF usará Helvetica y los símbolos fuera de cp1252 (≤ ≥ →) se escribirán como "
            "<= >= ->. Copie la carpeta fuentes/ del proyecto para evitarlo. El docx no se ve afectado.")
    return {"regular": "Helvetica", "negrita": "Helvetica-Bold", "latin1": True, "origen": None}


def _plano(texto, F: dict) -> str:
    """Texto para ``drawString``: sin caracteres de control; con fuentes estándar, lo que no existe en cp1252
    se sustituye por un equivalente legible (``≤`` -> ``<=``) o por ``?``."""
    t = _CONTROL.sub("", str(texto or ""))
    if F["latin1"]:
        for simbolo, equivalente in _SUSTITUCIONES_CP1252.items():
            t = t.replace(simbolo, equivalente)
        t = t.encode("cp1252", "replace").decode("cp1252")
    return t


def _txt(texto, F: dict) -> str:
    """Texto para ``Paragraph``: como ``_plano`` y además con XML escapado."""
    return _xml_escape(_plano(texto, F))


def _estilo(F: dict, pt: float, *, negrita: bool = False, color: str = COLOR_TEXTO,
            interlineado: float = INTERLINEADO, alineacion=TA_LEFT) -> ParagraphStyle:
    return ParagraphStyle(name="p", fontName=F["negrita"] if negrita else F["regular"], fontSize=pt,
                          leading=pt * interlineado, textColor=colors.HexColor(color), alignment=alineacion)


def _alto_parrafo(texto: str, estilo: ParagraphStyle, ancho_pt: float, F: dict) -> float:
    if not texto:
        return 0.0
    _w, h = Paragraph(_txt(texto, F), estilo).wrap(ancho_pt, 100_000)
    return h


def _lineas(texto: str, estilo: ParagraphStyle, ancho_pt: float, F: dict) -> int:
    """Líneas que ocupa el texto en un párrafo de ``ancho_pt`` (0 si está vacío)."""
    if not texto:
        return 0
    return max(1, int(round(_alto_parrafo(texto, estilo, ancho_pt, F) / estilo.leading)))


def _recortar_a_lineas(texto, estilo: ParagraphStyle, ancho_pt: float, lineas_max: int, F: dict) -> str:
    """Recorta el texto (con '…') hasta que ocupe como máximo ``lineas_max`` líneas de ``ancho_pt``.

    Se quitan palabras del final (búsqueda binaria sobre el número de palabras): se conserva todo lo que
    cabe, sin perder de golpe un 10 % del texto.
    """
    t = _limpiar(texto)

    def cabe(candidato: str) -> bool:
        return _alto_parrafo(candidato, estilo, ancho_pt, F) <= lineas_max * estilo.leading + 0.5

    if not t or cabe(t):
        return t
    palabras = t.split(" ")
    bajo, alto = 0, len(palabras) - 1        # bajo palabras siempre caben (0 = solo '…'); alto no cabe entero
    while alto - bajo > 1:
        medio = (bajo + alto) // 2
        if cabe(" ".join(palabras[:medio]).rstrip(" ,;:.") + "…"):
            bajo = medio
        else:
            alto = medio
    return " ".join(palabras[:bajo]).rstrip(" ,;:.") + "…"


def _recortar_ancho(texto, fuente: str, pt: float, ancho_max: float, F: dict) -> str:
    """Recorta a una sola línea de ``ancho_max`` puntos, medida con la fuente real.

    Se mide el texto tal como se dibujará (``_plano``) pero se devuelven los caracteres originales:
    el docx recibe siempre el texto Unicode íntegro.
    """
    t = _limpiar(texto)
    if pdfmetrics.stringWidth(_plano(t, F), fuente, pt) <= ancho_max:
        return t
    while len(t) > 1:
        t = t[:-1].rstrip(" ,;:.")
        if pdfmetrics.stringWidth(_plano(t + "…", F), fuente, pt) <= ancho_max:
            return t + "…"
    return "…"


@dataclass(frozen=True)
class _Celda:
    """Textos y medidas definitivos de una celda de paso (iguales en docx y pdf)."""

    meta: str
    titulo: str
    descripcion: str
    pt_titulo: float
    pt_desc: float
    img_w_cm: float
    img_h_cm: float
    aviso: str | None      # si hubo que recortar título o descripción, qué y cuánto


def _linea_paso(m: Momento, numero: int, L: Layout, F: dict, ancho_pt: float) -> str:
    """``"PASO N · SECCIÓN · mm:ss"``; la sección (lo único prescindible) se abrevia u omite si no cabe en una línea."""
    st_meta = _estilo(F, L.pt_meta)
    seccion = _limpiar(m.seccion).upper()
    tiempo = m.tiempo
    meta = " · ".join(p for p in (f"PASO {numero}", seccion, tiempo) if p)
    while seccion and _alto_parrafo(meta, st_meta, ancho_pt, F) > st_meta.leading + 0.5:
        nueva = _recortar(seccion, int(len(seccion) * 0.8)) if len(seccion) > 6 else ""
        seccion = nueva if len(nueva) < len(seccion) else ""
        meta = " · ".join(p for p in (f"PASO {numero}", seccion, tiempo) if p)
    return meta


def _maquetar_celda(m: Momento, numero: int, L: Layout, F: dict, ancho_pt: float) -> _Celda:
    """Decide textos, tamaños de fuente y caja de imagen de una celda para que el texto COMPLETO quepa.

    Orden: (1) con las fuentes de la maqueta, la imagen se reduce desde su tamaño de referencia hasta que
    quepan el título (≤ ``MAX_LINEAS_TITULO`` líneas) y la descripción completa (≤ ``L.lineas_desc_max``
    líneas); (2) si aun así no caben, se bajan 1 pt las fuentes; (3) solo entonces se recortan con '…' y se
    devuelve un aviso con el número de paso (nunca se recorta en silencio).
    """
    meta = _linea_paso(m, numero, L, F, ancho_pt)
    titulo0 = _limpiar(m.titulo) or "(sin título)"
    desc0 = _limpiar(m.descripcion)
    alto_celda_pt = (L.celda_h_cm - CELDA_MARGEN_SUP_CM) / 2.54 * 72
    img_ref_pt = L.img_h_cm / 2.54 * 72
    aviso = None
    for intento, (pt_t, pt_d) in enumerate(((L.pt_titulo, L.pt_desc), (L.pt_titulo - 1.0, L.pt_desc - 1.0))):
        st_t, st_d = _estilo(F, pt_t, negrita=True), _estilo(F, pt_d)
        titulo, descripcion = titulo0, desc0
        lt, ld = _lineas(titulo, st_t, ancho_pt, F), _lineas(descripcion, st_d, ancho_pt, F)
        if lt <= MAX_LINEAS_TITULO and ld <= L.lineas_desc_max:
            break
        if intento == 0:
            continue
        # último recurso: recortar (con la fuente ya reducida) y avisar
        recortes = []
        if lt > MAX_LINEAS_TITULO:
            titulo = _recortar_a_lineas(titulo0, st_t, ancho_pt, MAX_LINEAS_TITULO, F)
            recortes.append(f"título de {len(titulo0)} a {len(titulo)} caracteres")
        if ld > L.lineas_desc_max:
            descripcion = _recortar_a_lineas(desc0, st_d, ancho_pt, L.lineas_desc_max, F)
            recortes.append(f"descripción de {len(desc0)} a {len(descripcion)} caracteres")
        lt, ld = _lineas(titulo, st_t, ancho_pt, F), _lineas(descripcion, st_d, ancho_pt, F)
        aviso = (f"paso {numero}: no cabe el texto completo en la celda; se recortó " + " y ".join(recortes)
                 + f" (el texto íntegro está en {config.NOMBRE_JSON})")
    texto_pt = L.pt_meta * INTERLINEADO + lt * pt_t * INTERLINEADO + ld * pt_d * INTERLINEADO + PAD_TEXTO_PT
    img_h_pt = max(0.0, min(img_ref_pt, alto_celda_pt - texto_pt))
    img_h_cm = img_h_pt / 72 * 2.54
    img_w_cm = min(L.interior_w_cm, img_h_cm * (L.img_w_cm / L.img_h_cm if L.img_h_cm > 0 else L.ratio))
    return _Celda(meta=meta, titulo=titulo, descripcion=descripcion, pt_titulo=pt_t, pt_desc=pt_d,
                  img_w_cm=img_w_cm, img_h_cm=img_h_cm, aviso=aviso)


# ----------------------------------------------------------------------------- índice (paginado por medida)
@dataclass(frozen=True)
class _EntradaIndice:
    tipo: str        # "seccion" | "paso"
    texto: str       # título completo ("N. título"); ocupa ``lineas`` líneas
    tiempo: str
    numero: int      # número de paso (0 en las secciones)
    lineas: int
    alto: float      # pt (lineas x alto de línea)
    antes: float     # pt de espacio previo (se omite al inicio de una página)


def _ancho_texto_indice(L: Layout) -> float:
    """Ancho en pt de la columna de títulos del índice (una sola columna, también en A4 horizontal)."""
    return (L.util_w_cm - INDICE_ANCHO_DERECHA_CM) * cm


def _lineas_indice(texto: str, fuente: str, pt: float, ancho_pt: float, F: dict) -> int:
    """Líneas que ocupa una entrada del índice (``simpleSplit``: la misma partición con la que se dibuja)."""
    return max(1, len(simpleSplit(_plano(texto, F), fuente, pt, ancho_pt)))


def _entradas_indice(momentos: list, L: Layout, F: dict) -> list[_EntradaIndice]:
    """Entradas del índice con su alto medido: los títulos van completos (en varias líneas si hace falta)."""
    ancho_texto = _ancho_texto_indice(L)
    ancho_seccion = L.util_w_cm * cm
    entradas = []
    for nombre, pasos in _grupos_indice(momentos):
        lineas = _lineas_indice(nombre, F["negrita"], PT_INDICE_SECCION, ancho_seccion, F)
        entradas.append(_EntradaIndice("seccion", nombre, "", 0, lineas, lineas * ALTO_INDICE_SECCION_PT,
                                       ESPACIO_INDICE_SECCION_PT))
        for i, m in pasos:
            texto = f"{i}. {_limpiar(m.titulo)}"
            lineas = _lineas_indice(texto, F["regular"], PT_INDICE_PASO, ancho_texto, F)
            entradas.append(_EntradaIndice("paso", texto, m.tiempo, i, lineas, lineas * ALTO_INDICE_PASO_PT, 0.0))
    return entradas


def _paginar_indice(entradas: list[_EntradaIndice], L: Layout) -> list[list[_EntradaIndice]]:
    """Reparte las entradas en páginas por medida (una columna); una sección nunca queda sola al pie."""
    util_h_pt = (L.pag_h_cm - 2 * MARGEN_CM - CABECERA_CM - SEGURIDAD_CM) * cm

    def disponible(idx_pagina: int) -> float:
        return util_h_pt - (ALTO_TITULO_INDICE_PT if idx_pagina == 0 else 0.0)

    paginas: list[list[_EntradaIndice]] = []
    pagina: list[_EntradaIndice] = []
    restante = disponible(0)
    i = 0
    while i < len(entradas):
        e = entradas[i]
        antes = e.antes if pagina else 0.0
        necesario = antes + e.alto
        if e.tipo == "seccion" and i + 1 < len(entradas):
            necesario += entradas[i + 1].alto
        if necesario > restante and pagina:
            paginas.append(pagina)
            pagina = []
            restante = disponible(len(paginas))
            continue
        pagina.append(e)
        restante -= antes + e.alto
        i += 1
    if pagina:
        paginas.append(pagina)
    return paginas


def _pagina_del_paso(numero: int, n_indice: int, L: Layout) -> int:
    """Número de página (portada = 1) en que queda el paso ``numero``: portada + índice + páginas de pasos."""
    return 1 + n_indice + math.ceil(numero / L.por_pagina)


def _texto_derecha_indice(e: _EntradaIndice, n_indice: int, L: Layout) -> str:
    return f"{e.tiempo} · pág. {_pagina_del_paso(e.numero, n_indice, L)}"


@dataclass
class _Segmento:
    """Un segmento del anexo de transcripción ya medido (alto en pt del párrafo de texto)."""

    tiempo: str
    texto: str
    alto: float


def _estilo_transcripcion(F: dict) -> ParagraphStyle:
    return _estilo(F, PT_TRANSCRIPCION, color=COLOR_TEXTO)


def _ancho_texto_transcripcion(L: Layout) -> float:
    return (L.util_w_cm - TRANSCRIPCION_COL_CM) * cm


def _paginar_transcripcion(D: _Datos, L: Layout, F: dict) -> list[list[_Segmento]]:
    """Reparte los segmentos en páginas (la primera lleva el título); mismas páginas en docx y pdf."""
    if not D.transcripcion:
        return []
    estilo = _estilo_transcripcion(F)
    ancho = _ancho_texto_transcripcion(L)
    disponible = (L.pag_h_cm - 2 * MARGEN_CM - CABECERA_CM - SEGURIDAD_CM) * cm
    paginas: list[list[_Segmento]] = [[]]
    usado = ALTO_TITULO_INDICE_PT
    for tiempo, texto in D.transcripcion:
        alto = max(_alto_parrafo(texto, estilo, ancho, F), estilo.leading)
        if paginas[-1] and usado + alto + ESPACIO_SEGMENTO_PT > disponible:
            paginas.append([])
            usado = 0.0
        paginas[-1].append(_Segmento(tiempo=tiempo, texto=texto, alto=alto))
        usado += alto + ESPACIO_SEGMENTO_PT
    return paginas


def contar_paginas_indice(momentos: list, por_pagina="auto") -> int:
    """Páginas que ocupa el índice (misma paginación por medida que usan el PDF y el docx)."""
    momentos = list(momentos or [])
    if not momentos:
        return 0
    L = calcular_layout(len(momentos), por_pagina, ratio=_ratio_capturas(momentos))
    return len(_paginar_indice(_entradas_indice(momentos, L, _fuentes_pdf()), L))


# ============================================================================= DOCX
_SUCESORES_PPR = ("w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                  "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
                  "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                  "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                  "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def _rgb(color_hex: str) -> RGBColor:
    return RGBColor.from_string(color_hex.lstrip("#"))


def _parrafo(p, *, linea_pt: float | None = None, antes: float = 0, despues: float = 0, alineacion=None):
    """Espaciados deterministas; ``linea_pt`` fija el interlineado EXACTO (si es None, sencillo)."""
    pf = p.paragraph_format
    pf.space_before = Pt(antes)
    pf.space_after = Pt(despues)
    pf.line_spacing = Pt(linea_pt) if linea_pt else 1.0
    if alineacion is not None:
        p.alignment = alineacion
    return p


def _run(p, texto: str, pt: float, *, negrita: bool = False, cursiva: bool = False, color: str | None = None):
    r = p.add_run(_CONTROL.sub("", str(texto or "")))   # un carácter de control haría fallar todo el docx
    r.font.size = Pt(pt)
    r.bold = negrita
    r.italic = cursiva
    if color:
        r.font.color.rgb = _rgb(color)
    return r


def _espaciado_letras(run, pt: float) -> None:
    """Espaciado entre caracteres (w:spacing en vigésimos de punto)."""
    rPr = run._r.get_or_add_rPr()
    el = OxmlElement("w:spacing")
    el.set(qn("w:val"), str(int(round(pt * 20))))
    rPr.insert_element_before(el, "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u",
                              "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs",
                              "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")


def _borde_inferior(p, color: str = COLOR_LINEA) -> None:
    """Línea fina bajo el párrafo (w:pBdr/w:bottom)."""
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    abajo = OxmlElement("w:bottom")
    abajo.set(qn("w:val"), "single")
    abajo.set(qn("w:sz"), "4")          # octavos de punto
    abajo.set(qn("w:space"), "2")
    abajo.set(qn("w:color"), color.lstrip("#"))
    pBdr.append(abajo)
    pPr.insert_element_before(pBdr, "w:shd", *_SUCESORES_PPR)


def _sombreado(p, color: str = COLOR_CAJA) -> None:
    """Fondo de párrafo (w:shd)."""
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color.lstrip("#"))
    pPr.insert_element_before(shd, *_SUCESORES_PPR)


def _tab_derecha(p, pos_cm: float, puntos: bool = False) -> None:
    p.paragraph_format.tab_stops.add_tab_stop(Cm(pos_cm), WD_TAB_ALIGNMENT.RIGHT,
                                              WD_TAB_LEADER.DOTS if puntos else WD_TAB_LEADER.SPACES)


def _campo(p, instruccion: str, pt: float, color: str) -> None:
    """Campo simple de Word (PAGE, NUMPAGES) con el formato indicado."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instruccion)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    col = OxmlElement("w:color")
    col.set(qn("w:val"), color.lstrip("#"))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(round(pt * 2))))
    rPr.append(col)
    rPr.append(sz)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


def _tabla_margenes_celda(tabla, sup_cm: float, lat_cm: float, inf_cm: float = 0.0) -> None:
    """Márgenes internos de celda (w:tblCellMar), insertado en orden de esquema (antes de tblLook)."""
    tblPr = tabla._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for lado, val in (("top", sup_cm), ("left", lat_cm), ("bottom", inf_cm), ("right", lat_cm)):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:w"), str(int(round(val / 2.54 * 1440))))   # twips
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    look = tblPr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(mar)
    else:
        tblPr.append(mar)


def _configurar_docx(doc, D: _Datos, L: Layout) -> None:
    """Estilo base, tamaño/orientación de página, márgenes y pie con número de página."""
    normal = doc.styles["Normal"]
    normal.font.name = config.FUENTE_DOCX
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), config.FUENTE_DOCX)
    normal.font.size = Pt(10)
    normal.font.color.rgb = _rgb(COLOR_TEXTO)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE if L.horizontal else WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(L.pag_w_cm), Cm(L.pag_h_cm)   # python-docx no intercambia solos
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(MARGEN_CM)
    sec.header_distance = sec.footer_distance = Cm(0.5)
    sec.different_first_page_header_footer = True      # la portada no lleva número de página
    pie = _parrafo(sec.footer.paragraphs[0], linea_pt=PT_CABECERA * INTERLINEADO,
                   alineacion=WD_ALIGN_PARAGRAPH.RIGHT)
    _run(pie, "Página ", PT_CABECERA, color=COLOR_GRIS)
    _campo(pie, "PAGE", PT_CABECERA, COLOR_GRIS)
    _run(pie, " de ", PT_CABECERA, color=COLOR_GRIS)
    _campo(pie, "NUMPAGES", PT_CABECERA, COLOR_GRIS)
    doc.core_properties.title = D.titulo
    doc.core_properties.author = "resumen_videos"
    doc.core_properties.subject = f"{len(D.momentos)} pasos"


def _portada_docx(doc, D: _Datos, L: Layout) -> None:
    sangria = Cm(1.5)
    p = _parrafo(doc.add_paragraph(), antes=L.pag_h_cm * (0.10 if L.horizontal else 0.13) / 2.54 * 72)
    p.paragraph_format.left_indent = p.paragraph_format.right_indent = sangria
    _espaciado_letras(_run(p, "MANUAL DE PROCEDIMIENTO", 9, negrita=True, color=COLOR_GRIS), 1.2)

    p = _parrafo(doc.add_paragraph(), antes=10, despues=8)
    p.paragraph_format.left_indent = p.paragraph_format.right_indent = sangria
    _run(p, D.titulo, 26, negrita=True, color=COLOR_TITULO)
    _borde_inferior(p)

    if D.resumen:
        p = _parrafo(doc.add_paragraph(), antes=12, despues=6)
        p.paragraph_format.left_indent = p.paragraph_format.right_indent = sangria
        _run(p, D.resumen, 11.5, color=COLOR_TEXTO)

    for i, (etiqueta, valor) in enumerate(D.filas_portada()):
        p = _parrafo(doc.add_paragraph(), antes=14 if i == 0 else 3)
        p.paragraph_format.left_indent = p.paragraph_format.right_indent = sangria
        p.paragraph_format.tab_stops.add_tab_stop(sangria + Cm(3.2), WD_TAB_ALIGNMENT.LEFT)
        _run(p, etiqueta.upper() + "\t", 8, color=COLOR_GRIS)
        _run(p, valor, 10.5, color=COLOR_TITULO)

    p = _parrafo(doc.add_paragraph(), antes=40)
    p.paragraph_format.left_indent = p.paragraph_format.right_indent = sangria
    _run(p, D.pie, 8, color=COLOR_GRIS)


def _cabecera_docx(doc, L: Layout, izquierda: str, derecha: str):
    """Párrafo de cabecera con salto de página ANTES y línea fina debajo."""
    cab = _parrafo(doc.add_paragraph(), linea_pt=L.pt_cabecera * INTERLINEADO, despues=4)
    cab.paragraph_format.page_break_before = True
    _tab_derecha(cab, L.util_w_cm)
    _run(cab, _recortar(izquierda, 90), L.pt_cabecera, color=COLOR_GRIS)
    _run(cab, "\t" + derecha, L.pt_cabecera, color=COLOR_GRIS)
    _borde_inferior(cab)
    return cab


def _indice_docx(doc, D: _Datos, L: Layout, paginas_indice: list) -> None:
    """Índice paginado igual que en el PDF: una tabla sin bordes por página (títulos | mm:ss · pág.), filas EXACTAS.

    Cada página de índice empieza con la cabecera (salto de página antes) y las filas tienen el alto medido con
    las fuentes del PDF (Calibri es más estrecha: nunca necesita más líneas), así docx y pdf coinciden página a
    página y el número de página de cada paso es el mismo en ambos.
    """
    n_indice = len(paginas_indice)
    ancho_der_cm = INDICE_ANCHO_DERECHA_CM
    ancho_izq_cm = L.util_w_cm - ancho_der_cm
    for k, entradas in enumerate(paginas_indice):
        _cabecera_docx(doc, L, D.titulo, "Índice" if n_indice == 1 else f"Índice ({k + 1} de {n_indice})")
        if k == 0:
            p = _parrafo(doc.add_paragraph(), linea_pt=PT_INDICE_TITULO * INTERLINEADO,
                         despues=ALTO_TITULO_INDICE_PT - PT_INDICE_TITULO * INTERLINEADO)
            _run(p, "Índice", PT_INDICE_TITULO, negrita=True, color=COLOR_TITULO)
        tabla = doc.add_table(rows=len(entradas), cols=2)
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        tabla.autofit = False
        _tabla_margenes_celda(tabla, 0.0, 0.0)
        for j, e in enumerate(entradas):
            antes = e.antes if j else 0.0
            fila = tabla.rows[j]
            fila.height = Pt(antes + e.alto)
            fila.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            izq, der = fila.cells
            izq.width, der.width = Cm(ancho_izq_cm), Cm(ancho_der_cm)
            if e.tipo == "seccion":
                p = _parrafo(izq.paragraphs[0], linea_pt=ALTO_INDICE_SECCION_PT, antes=antes)
                _run(p, e.texto, PT_INDICE_SECCION, negrita=True, color=COLOR_TITULO)
                continue
            p = _parrafo(izq.paragraphs[0], linea_pt=ALTO_INDICE_PASO_PT, antes=antes)
            _run(p, e.texto, PT_INDICE_PASO, color=COLOR_TEXTO)
            # tiempo y página alineados con la ÚLTIMA línea del título, con puntos de guía
            p = _parrafo(der.paragraphs[0], linea_pt=ALTO_INDICE_PASO_PT,
                         antes=antes + (e.lineas - 1) * ALTO_INDICE_PASO_PT, alineacion=WD_ALIGN_PARAGRAPH.RIGHT)
            _tab_derecha(p, ancho_der_cm - 0.1, puntos=True)
            _run(p, "\t" + _texto_derecha_indice(e, n_indice, L), PT_INDICE_PASO, color=COLOR_GRIS)


def _caja_sin_captura_docx(celda, C: _Celda, L: Layout) -> None:
    """Caja gris del alto de la imagen con '[sin captura]' centrado (tres párrafos sombreados)."""
    img_h_pt = C.img_h_cm / 2.54 * 72
    alto_texto = C.pt_desc * INTERLINEADO
    relleno = max(1.0, (img_h_pt - alto_texto) / 2)
    partes = ((celda.paragraphs[0], relleno, ""), (celda.add_paragraph(), alto_texto, "[sin captura]"),
              (celda.add_paragraph(), relleno, ""))
    for k, (p, alto, texto) in enumerate(partes):
        _parrafo(p, linea_pt=alto, despues=ESPACIO_TRAS_IMAGEN_PT if k == 2 else 0,
                 alineacion=WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.right_indent = Cm(max(0.0, L.interior_w_cm - C.img_w_cm))   # caja del ancho de la imagen
        _sombreado(p)
        _run(p, texto, C.pt_desc if texto else 1, color=COLOR_GRIS)


def _marca_parrafo_pequena(p, pt: float = 2.0) -> None:
    """Tamaño de la marca de párrafo (w:pPr/w:rPr/w:sz): en Word también fija el alto de la última línea."""
    pPr = p._p.get_or_add_pPr()
    rPr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(round(pt * 2))))
    rPr.append(sz)
    pPr.append(rPr)   # w:rPr va al final de w:pPr (solo le siguen sectPr y pPrChange)


def _celda_momento_docx(celda, m: Momento, numero: int, L: Layout, F: dict, log: Callable[[str], None],
                        avisos: list) -> None:
    ancho_pt = (L.celda_w_cm - 2 * CELDA_MARGEN_LAT_CM) * cm
    C = _maquetar_celda(m, numero, L, F, ancho_pt)
    if C.aviso:
        avisos.append(C.aviso)
    ruta = m.captura_para_documento
    tam = _tamano_imagen(ruta)
    insertada = False
    if tam and C.img_h_cm > 0:
        # imagen (párrafo 0 de la celda; sin interlineado exacto o Word recorta la imagen)
        p_img = _parrafo(celda.paragraphs[0], despues=ESPACIO_TRAS_IMAGEN_PT, alineacion=WD_ALIGN_PARAGRAPH.LEFT)
        _marca_parrafo_pequena(p_img)   # la marca de párrafo hereda 10 pt y agrandaría la línea de la imagen
        run = p_img.add_run()
        run.font.size = Pt(2)   # minimiza el descendente de la línea que contiene la imagen
        try:
            if tam[0] / tam[1] >= C.img_w_cm / C.img_h_cm:
                run.add_picture(str(ruta), width=Cm(C.img_w_cm))
            else:
                run.add_picture(str(ruta), height=Cm(C.img_h_cm))
            insertada = True
        except (OSError, InvalidImageStreamError) as exc:
            log(f"  aviso: captura ilegible para el docx ({Path(str(ruta)).name}): {exc}")
            run._r.getparent().remove(run._r)
    if not insertada:
        _caja_sin_captura_docx(celda, C, L)
    p = _parrafo(celda.add_paragraph(), linea_pt=L.pt_meta * INTERLINEADO, despues=ESPACIO_ENTRE_PARRAFOS_PT)
    _run(p, C.meta, L.pt_meta, color=COLOR_GRIS)
    p = _parrafo(celda.add_paragraph(), linea_pt=C.pt_titulo * INTERLINEADO, despues=ESPACIO_ENTRE_PARRAFOS_PT)
    _run(p, C.titulo, C.pt_titulo, negrita=True, color=COLOR_TITULO)
    p = _parrafo(celda.add_paragraph(), linea_pt=C.pt_desc * INTERLINEADO)
    _run(p, C.descripcion, C.pt_desc, color=COLOR_TEXTO)


def _texto_rango(ini: int, fin: int, n: int) -> str:
    return f"Paso {ini} de {n}" if ini == fin else f"Pasos {ini}–{fin} de {n}"


def _pagina_contenido_docx(doc, pagina: list, k: int, D: _Datos, L: Layout, F: dict,
                           log: Callable[[str], None], avisos: list) -> None:
    ini = k * L.por_pagina + 1
    fin = ini + len(pagina) - 1
    _cabecera_docx(doc, L, D.titulo, _texto_rango(ini, fin, L.n))
    # rejilla: tabla sin estilo (sin bordes), ancho fijo por celda, filas de alto EXACTO
    tabla = doc.add_table(rows=L.filas, cols=L.cols)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = False
    _tabla_margenes_celda(tabla, CELDA_MARGEN_SUP_CM, CELDA_MARGEN_LAT_CM)
    for f in range(L.filas):
        fila = tabla.rows[f]
        fila.height = Cm(L.celda_h_cm)
        fila.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY   # Word recorta en vez de crecer: paginación fija
        for c in range(L.cols):
            celda = fila.cells[c]
            celda.width = Cm(L.celda_w_cm)              # column.width no basta
            celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            idx = f * L.cols + c
            if idx < len(pagina):
                _celda_momento_docx(celda, pagina[idx], ini + idx, L, F, log, avisos)


def _transcripcion_docx(doc, D: _Datos, L: Layout, F: dict) -> None:
    """Anexo con la transcripción: una página por grupo de ``_paginar_transcripcion`` (igual que el PDF)."""
    paginas = _paginar_transcripcion(D, L, F)
    n = len(paginas)
    col = Cm(TRANSCRIPCION_COL_CM)
    for k, segmentos in enumerate(paginas):
        _cabecera_docx(doc, L, D.titulo, "Transcripción" if n == 1 else f"Transcripción ({k + 1} de {n})")
        if k == 0:
            p = _parrafo(doc.add_paragraph(), linea_pt=PT_TRANSCRIPCION_TITULO * INTERLINEADO,
                         despues=ALTO_TITULO_INDICE_PT - PT_TRANSCRIPCION_TITULO * INTERLINEADO)
            _run(p, "Transcripción del audio", PT_TRANSCRIPCION_TITULO, negrita=True, color=COLOR_TITULO)
        for s in segmentos:
            p = _parrafo(doc.add_paragraph(), linea_pt=PT_TRANSCRIPCION * INTERLINEADO, despues=ESPACIO_SEGMENTO_PT)
            p.paragraph_format.left_indent = col
            p.paragraph_format.first_line_indent = -col
            p.paragraph_format.tab_stops.add_tab_stop(col, WD_TAB_ALIGNMENT.LEFT)
            _run(p, s.tiempo + "\t", PT_TRANSCRIPCION, color=COLOR_GRIS)
            _run(p, s.texto, PT_TRANSCRIPCION, color=COLOR_TEXTO)


def generar_docx(nombre_video: str, momentos: list, carpeta_salida: Path, *, titulo: str | None = None,
                 resumen: str | None = None, fecha: str | None = None, duracion: float | None = None,
                 modo: str = "", modelo: str = "", por_pagina="auto", incluir_indice: bool = True,
                 transcripcion: list | None = None,
                 log: Callable[[str], None] = print, avisos: list | None = None) -> Path:
    """Escribe ``<carpeta_salida>/<nombre_video>.docx`` y devuelve su ruta.  ``momentos`` vacío -> ValueError.

    Si se pasa ``avisos`` (lista), se le añaden los recortes de texto que hubo que hacer (uno por paso).
    ``transcripcion`` (``[{"inicio": s, "fin": s, "texto": str}]``) se añade como anexo al final.
    """
    avisos = avisos if avisos is not None else []
    D, L, F = _preparar(nombre_video, momentos, titulo, resumen, fecha, duracion, modo, modelo, por_pagina, log,
                        transcripcion=transcripcion)
    doc = Document()
    _configurar_docx(doc, D, L)
    _portada_docx(doc, D, L)
    if incluir_indice:
        _indice_docx(doc, D, L, _paginar_indice(_entradas_indice(D.momentos, L, F), L))
    for k, pagina in enumerate(_paginar(D.momentos, L.por_pagina)):
        _pagina_contenido_docx(doc, pagina, k, D, L, F, log, avisos)
    _transcripcion_docx(doc, D, L, F)
    # Word exige un párrafo tras la última tabla; de 1 pt para que no genere una página más
    fin = _parrafo(doc.add_paragraph(), linea_pt=1)
    _run(fin, "", 1)
    ruta = Path(carpeta_salida) / f"{_nombre_archivo_seguro(nombre_video)}.docx"
    ruta.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(ruta))
    return ruta


# ============================================================================= PDF
def _cabecera_pdf(c, L: Layout, F: dict, izquierda: str, derecha: str) -> None:
    """Texto gris a ambos lados y línea fina debajo."""
    W, H = L.pag_w_cm * cm, L.pag_h_cm * cm
    mg = MARGEN_CM * cm
    y = H - mg - L.pt_cabecera
    derecha = _plano(derecha, F)
    ancho_der = pdfmetrics.stringWidth(derecha, F["regular"], L.pt_cabecera)
    izquierda = _recortar_ancho(izquierda, F["regular"], L.pt_cabecera, W - 2 * mg - ancho_der - 20, F)
    c.setFont(F["regular"], L.pt_cabecera)
    c.setFillColor(colors.HexColor(COLOR_GRIS))
    c.drawString(mg, y, izquierda)
    c.drawRightString(W - mg, y, derecha)
    c.setStrokeColor(colors.HexColor(COLOR_LINEA))
    c.setLineWidth(0.5)
    c.line(mg, y - 4, W - mg, y - 4)


def _portada_pdf(c, D: _Datos, L: Layout, F: dict) -> None:
    W, H = L.pag_w_cm * cm, L.pag_h_cm * cm
    x = 2.5 * cm
    ancho = W - 2 * x
    y = H - (3.2 if L.horizontal else 4.5) * cm
    c.setFillColor(colors.HexColor(COLOR_GRIS))
    encabezado = c.beginText(x, y)
    encabezado.setFont(F["negrita"], 9)
    encabezado.setCharSpace(1.2)
    encabezado.textOut("MANUAL DE PROCEDIMIENTO")
    encabezado.setCharSpace(0)   # el espaciado (Tc) persiste en el estado gráfico de la página
    c.drawText(encabezado)
    y -= 20
    p = Paragraph(_txt(D.titulo, F), _estilo(F, 26, negrita=True, color=COLOR_TITULO))
    _w, h = p.wrap(ancho, H)
    p.drawOn(c, x, y - h)
    y -= h + 12
    c.setStrokeColor(colors.HexColor(COLOR_LINEA))
    c.setLineWidth(0.8)
    c.line(x, y, x + ancho, y)
    y -= 22
    if D.resumen:
        p = Paragraph(_txt(D.resumen, F), _estilo(F, 11.5, interlineado=1.4))
        _w, h = p.wrap(ancho * 0.9, H)
        p.drawOn(c, x, y - h)
        y -= h + 26
    for etiqueta, valor in D.filas_portada():
        c.setFont(F["regular"], 8)
        c.setFillColor(colors.HexColor(COLOR_GRIS))
        c.drawString(x, y, _plano(etiqueta.upper(), F))
        c.setFont(F["regular"], 10.5)
        c.setFillColor(colors.HexColor(COLOR_TITULO))
        c.drawString(x + 3.2 * cm, y, _recortar_ancho(valor, F["regular"], 10.5, ancho - 3.2 * cm, F))
        y -= 17
    c.setFont(F["regular"], 8)
    c.setFillColor(colors.HexColor(COLOR_GRIS))
    c.drawString(x, 1.6 * cm, _recortar_ancho(D.pie, F["regular"], 8, ancho, F))


def _pagina_indice_pdf(c, entradas: list, k: int, num_pagina: int, total: int, n_indice: int, D: _Datos,
                       L: Layout, F: dict) -> None:
    """Una página del índice: una sola columna, títulos completos (varias líneas) y 'mm:ss · pág. N' a la derecha."""
    H = L.pag_h_cm * cm
    mg = MARGEN_CM * cm
    _cabecera_pdf(c, L, F, D.titulo, f"Índice · Página {num_pagina} de {total}")
    y_top = H - mg - CABECERA_CM * cm
    if k == 0:
        c.setFont(F["negrita"], PT_INDICE_TITULO)
        c.setFillColor(colors.HexColor(COLOR_TITULO))
        c.drawString(mg, y_top - PT_INDICE_TITULO * INTERLINEADO + 4, "Índice")
        y_top -= ALTO_TITULO_INDICE_PT
    ancho_col = L.util_w_cm * cm
    ancho_texto = _ancho_texto_indice(L)
    punto_w = pdfmetrics.stringWidth(".", F["regular"], PT_INDICE_PASO)
    x = mg
    y = y_top
    for j, e in enumerate(entradas):
        y -= (e.antes if j else 0.0) + e.alto
        alto_linea = e.alto / e.lineas
        if e.tipo == "seccion":
            c.setFont(F["negrita"], PT_INDICE_SECCION)
            c.setFillColor(colors.HexColor(COLOR_TITULO))
            lineas = simpleSplit(_plano(e.texto, F), F["negrita"], PT_INDICE_SECCION, ancho_col)
            for li, linea in enumerate(lineas[:e.lineas]):
                c.drawString(x, y + e.alto - (li + 1) * alto_linea + 0.3 * alto_linea, linea)
            continue
        lineas = simpleSplit(_plano(e.texto, F), F["regular"], PT_INDICE_PASO, ancho_texto)[:e.lineas]
        c.setFont(F["regular"], PT_INDICE_PASO)
        c.setFillColor(colors.HexColor(COLOR_TEXTO))
        for li, linea in enumerate(lineas):
            c.drawString(x, y + e.alto - (li + 1) * alto_linea + 0.3 * alto_linea, linea)
        # puntos de guía, tiempo y página en la última línea del título
        base = y + 0.3 * alto_linea
        derecha = _plano(_texto_derecha_indice(e, n_indice, L), F)
        derecha_w = pdfmetrics.stringWidth(derecha, F["regular"], PT_INDICE_PASO)
        x_fin_texto = x + pdfmetrics.stringWidth(lineas[-1] if lineas else "", F["regular"], PT_INDICE_PASO) + 3
        x_derecha = x + ancho_col - derecha_w
        puntos = int((x_derecha - 3 - x_fin_texto) // punto_w)
        c.setFillColor(colors.HexColor(COLOR_GRIS))
        if puntos > 0:
            c.drawString(x_derecha - 3 - puntos * punto_w, base, "." * puntos)
        c.drawString(x_derecha, base, derecha)


def _imagen_pdf(c, ruta, x: float, y: float, w: float, h: float, F: dict, log: Callable[[str], None]) -> None:
    """Captura ajustada por aspecto (escala mínima), pegada al borde superior izquierdo de la caja."""
    tam = _tamano_imagen(ruta)
    if tam:
        try:
            esc = min(w / tam[0], h / tam[1])
            dw, dh = tam[0] * esc, tam[1] * esc
            c.drawImage(ImageReader(str(ruta)), x, y + (h - dh), dw, dh)
            return
        except Exception as exc:  # noqa: BLE001 - reportlab lanza tipos heterogéneos por imágenes dañadas
            log(f"  aviso: captura ilegible para el pdf ({Path(str(ruta)).name}): {exc}")
    c.setFillColor(colors.HexColor(COLOR_CAJA))
    c.rect(x, y, w, h, stroke=0, fill=1)
    c.setFillColor(colors.HexColor(COLOR_GRIS))
    c.setFont(F["regular"], 9)
    c.drawCentredString(x + w / 2, y + h / 2 - 3, "[sin captura]")


def _celda_pdf(c, m: Momento, numero: int, x0: float, y1: float, L: Layout, F: dict,
               log: Callable[[str], None], avisos: list) -> None:
    cw, ch = L.celda_w_cm * cm, L.celda_h_cm * cm
    lat = CELDA_MARGEN_LAT_CM * cm
    tx, tw = x0 + lat, cw - 2 * lat
    C = _maquetar_celda(m, numero, L, F, tw)
    if C.aviso:
        avisos.append(C.aviso)
    iw, ih = C.img_w_cm * cm, C.img_h_cm * cm
    # imagen y textos alineados al borde izquierdo interior de la celda (como en el docx)
    ix, iy = x0 + lat, y1 - CELDA_MARGEN_SUP_CM * cm - ih
    _imagen_pdf(c, m.captura_para_documento, ix, iy, iw, ih, F, log)
    y = iy - 3
    for texto, estilo in ((C.meta, _estilo(F, L.pt_meta, color=COLOR_GRIS)),
                          (C.titulo, _estilo(F, C.pt_titulo, negrita=True, color=COLOR_TITULO)),
                          (C.descripcion, _estilo(F, C.pt_desc))):
        if not texto:
            continue
        p = Paragraph(_txt(texto, F), estilo)
        _w, h = p.wrap(tw, ch)
        p.drawOn(c, tx, y - h)
        y -= h + 1


def _pagina_contenido_pdf(c, pagina: list, k: int, num_pagina: int, total: int, D: _Datos, L: Layout,
                          F: dict, log: Callable[[str], None], avisos: list) -> None:
    H = L.pag_h_cm * cm
    mg = MARGEN_CM * cm
    ini = k * L.por_pagina + 1
    fin = ini + len(pagina) - 1
    _cabecera_pdf(c, L, F, D.titulo, f"{_texto_rango(ini, fin, L.n)} · Página {num_pagina} de {total}")
    y_top = H - mg - CABECERA_CM * cm
    for idx, m in enumerate(pagina):
        f, col = divmod(idx, L.cols)
        x0 = mg + col * (L.celda_w_cm + GAP_X_CM) * cm
        y1 = y_top - f * (L.celda_h_cm + GAP_Y_CM) * cm
        _celda_pdf(c, m, ini + idx, x0, y1, L, F, log, avisos)


def _pagina_transcripcion_pdf(c, segmentos: list, k: int, n: int, num_pagina: int, total: int, D: _Datos,
                              L: Layout, F: dict) -> None:
    W, H = L.pag_w_cm * cm, L.pag_h_cm * cm
    mg = MARGEN_CM * cm
    etiqueta = "Transcripción" if n == 1 else f"Transcripción ({k + 1} de {n})"
    _cabecera_pdf(c, L, F, D.titulo, f"{etiqueta} · Página {num_pagina} de {total}")
    y = H - mg - CABECERA_CM * cm
    if k == 0:
        c.setFont(F["negrita"], PT_TRANSCRIPCION_TITULO)
        c.setFillColor(colors.HexColor(COLOR_TITULO))
        c.drawString(mg, y - PT_TRANSCRIPCION_TITULO * INTERLINEADO + 4, _plano("Transcripción del audio", F))
        y -= ALTO_TITULO_INDICE_PT
    estilo = _estilo_transcripcion(F)
    ancho = _ancho_texto_transcripcion(L)
    for s in segmentos:
        parrafo = Paragraph(_txt(s.texto, F), estilo)
        _w, h = parrafo.wrap(ancho, 100_000)
        c.setFont(F["regular"], PT_TRANSCRIPCION)
        c.setFillColor(colors.HexColor(COLOR_GRIS))
        c.drawString(mg, y - estilo.leading + 3, _plano(s.tiempo, F))
        parrafo.drawOn(c, mg + TRANSCRIPCION_COL_CM * cm, y - h)
        y -= max(h, s.alto) + ESPACIO_SEGMENTO_PT


def generar_pdf(nombre_video: str, momentos: list, carpeta_salida: Path, *, titulo: str | None = None,
                resumen: str | None = None, fecha: str | None = None, duracion: float | None = None,
                modo: str = "", modelo: str = "", por_pagina="auto", incluir_indice: bool = True,
                transcripcion: list | None = None,
                log: Callable[[str], None] = print, avisos: list | None = None) -> Path:
    """Escribe ``<carpeta_salida>/<nombre_video>.pdf`` y devuelve su ruta.  ``momentos`` vacío -> ValueError.

    Si se pasa ``avisos`` (lista), se le añaden los recortes de texto que hubo que hacer (uno por paso).
    ``transcripcion`` (``[{"inicio": s, "fin": s, "texto": str}]``) se añade como anexo al final.
    """
    avisos = avisos if avisos is not None else []
    D, L, F = _preparar(nombre_video, momentos, titulo, resumen, fecha, duracion, modo, modelo, por_pagina, log,
                        transcripcion=transcripcion)
    paginas_indice = _paginar_indice(_entradas_indice(D.momentos, L, F), L) if incluir_indice else []
    paginas_transcripcion = _paginar_transcripcion(D, L, F)
    total = 1 + len(paginas_indice) + L.paginas_contenido + len(paginas_transcripcion)
    ruta = Path(carpeta_salida) / f"{_nombre_archivo_seguro(nombre_video)}.pdf"
    ruta.parent.mkdir(parents=True, exist_ok=True)
    c = rl_canvas.Canvas(str(ruta), pagesize=(L.pag_w_cm * cm, L.pag_h_cm * cm))
    c.setTitle(D.titulo)
    c.setAuthor("resumen_videos")
    c.setSubject(f"{L.n} pasos")
    _portada_pdf(c, D, L, F)
    c.showPage()
    numero = 2
    for k, entradas in enumerate(paginas_indice):
        _pagina_indice_pdf(c, entradas, k, numero, total, len(paginas_indice), D, L, F)
        c.showPage()
        numero += 1
    for k, pagina in enumerate(_paginar(D.momentos, L.por_pagina)):
        _pagina_contenido_pdf(c, pagina, k, numero, total, D, L, F, log, avisos)
        c.showPage()
        numero += 1
    for k, segmentos in enumerate(paginas_transcripcion):
        _pagina_transcripcion_pdf(c, segmentos, k, len(paginas_transcripcion), numero, total, D, L, F)
        c.showPage()
        numero += 1
    c.save()
    return ruta


# ============================================================================= API principal
def contar_paginas_pdf(ruta: Path) -> int:
    """Número de páginas del PDF (pymupdf, import perezoso); -1 si no se puede contar."""
    try:
        import pymupdf  # noqa: PLC0415 - perezoso: la generación no depende de pymupdf
        with pymupdf.open(str(ruta)) as doc:
            return doc.page_count
    except Exception:  # noqa: BLE001 - ImportError o cualquier error de lectura: -1 según la especificación
        return -1


def generar_documentos(nombre_video: str, momentos: list, carpeta_salida: Path, *, titulo: str | None = None,
                       resumen: str | None = None, fecha: str | None = None, duracion: float | None = None,
                       modo: str = "", modelo: str = "", por_pagina="auto", incluir_indice: bool = True,
                       transcripcion: list | None = None,
                       log: Callable[[str], None] = print, avisos: list | None = None) -> tuple[Path, Path, int]:
    """Genera ``<carpeta>/<nombre_video>.docx`` y ``.pdf``; devuelve ``(docx, pdf, paginas_pdf)``.

    ``paginas_pdf`` se cuenta con pymupdf (-1 si no está disponible).  ``momentos`` vacío -> ValueError.
    Los textos nunca se recortan en silencio: si en algún paso hubo que recortar (tras reducir la imagen y
    bajar la fuente), se escribe un aviso con el número de paso por ``log`` y se añade a ``avisos`` si se
    pasa una lista.
    """
    momentos = list(momentos or [])
    if not momentos:
        raise ValueError("no hay momentos: no se puede generar el documento")
    comunes = dict(titulo=titulo, resumen=resumen, fecha=fecha, duracion=duracion, modo=modo, modelo=modelo,
                   por_pagina=por_pagina, incluir_indice=incluir_indice, transcripcion=transcripcion, log=log)
    L = calcular_layout(len(momentos), por_pagina, ratio=_ratio_capturas(momentos))
    log(f"  documentos: {len(momentos)} pasos, {L.por_pagina} por página ({L.descripcion}"
        f"{'; capturas verticales' if L.vertical else ''})")
    recortes: list = []
    ruta_docx = generar_docx(nombre_video, momentos, carpeta_salida, avisos=recortes, **comunes)
    log(f"  docx: {ruta_docx.name}")
    ruta_pdf = generar_pdf(nombre_video, momentos, carpeta_salida, **comunes)   # mismas celdas: mismos recortes
    paginas = contar_paginas_pdf(ruta_pdf)
    log(f"  pdf: {ruta_pdf.name} ({paginas if paginas >= 0 else '?'} páginas)")
    for recorte in recortes:
        log(f"  aviso: {recorte}")
    if avisos is not None:
        avisos.extend(recortes)
    return ruta_docx, ruta_pdf, paginas
